import os
import re
import secrets
from pathlib import Path

import openpyxl
from werkzeug.security import generate_password_hash

from app import db
from app.models import Member, User
from app.utils import clean_text, workbook_path

INITIAL_ACCOUNTS = (
    ("admin", "ADMIN_PASSWORD", "Registry Administrator", "admin"),
    ("staff", "STAFF_PASSWORD", "Registry Staff", "staff"),
)


def ensure_users():
    created = []
    for username, env_key, full_name, role in INITIAL_ACCOUNTS:
        if User.query.filter_by(username=username).first():
            continue
        password = os.environ.get(env_key) or secrets.token_urlsafe(12)
        db.session.add(
            User(
                username=username,
                password_hash=generate_password_hash(password),
                full_name=full_name,
                role=role,
                must_change_password=True,
            )
        )
        created.append((username, password, full_name, role))
    if created:
        db.session.commit()
        _write_account_notice(created)
    return bool(created)


def _write_account_notice(created):
    path = Path(__file__).resolve().parent.parent / "data" / "initial_accounts.txt"
    path.parent.mkdir(parents=True, exist_ok=True)
    lines = [
        "Municipality of Jagna",
        "Motorela Registry — initial sign-in",
        "",
        "Replace these passwords after the first sign-in. The registry asks for that",
        "before it opens the desk. This file is not committed to git.",
        "",
    ]
    for username, password, full_name, role in created:
        lines.extend(
            [
                f"{full_name} ({role})",
                f"  username: {username}",
                f"  password: {password}",
                "",
            ]
        )
    path.write_text("\n".join(lines), encoding="utf-8")


def import_members(*, replace=False):
    path = workbook_path()
    if path is None:
        raise FileNotFoundError(
            "The motorela workbook was not found. Place the Excel file in data/ or set EXCEL_PATH."
        )
    if replace:
        Member.query.delete()
        db.session.commit()
    elif Member.query.count():
        return 0

    workbook = openpyxl.load_workbook(path, data_only=True)
    try:
        sheet = workbook.active
        header = next(sheet.iter_rows(min_row=4, max_row=4, values_only=True), None)
        year_row = next(sheet.iter_rows(min_row=2, max_row=2, values_only=True), None)
        if not header or clean_text(header[0]).lower() != "bodynumber":
            raise RuntimeError(f"{path.name} is missing the BodyNumber header on row 4.")
        record_year = 2026
        if year_row and year_row[0]:
            found = re.search(r"(20\d{2})", str(year_row[0]))
            if found:
                record_year = int(found.group(1))
        members = []
        seen = set()
        for row in sheet.iter_rows(min_row=5, values_only=True):
            if not row or not clean_text(row[0]):
                continue
            body = clean_text(row[0], upper=True)
            if body in seen:
                continue
            seen.add(body)
            owner = clean_text(row[1] if len(row) > 1 else "", upper=True)
            address = clean_text(row[2] if len(row) > 2 else "")
            association = clean_text(row[3] if len(row) > 3 else "", upper=True) or None
            franchise = clean_text(row[4] if len(row) > 4 else "", upper=True)
            driver = clean_text(row[5] if len(row) > 5 else "", upper=True) or "OWNER"
            members.append(
                Member(
                    body_number=body,
                    owner_name=owner,
                    address=address,
                    association=association,
                    franchise_number=franchise,
                    driver_name=driver,
                    driver_is_owner=driver == "OWNER" or driver == owner,
                    status="Active",
                    record_year=record_year,
                )
            )
        if not members:
            raise RuntimeError(f"{path.name} did not contain any body numbers.")
        db.session.add_all(members)
        db.session.commit()
        return len(members)
    except Exception:
        db.session.rollback()
        raise
    finally:
        workbook.close()


def ensure_seeded():
    db.create_all()
    ensure_users()
    imported = import_members()
    if imported:
        print(f"Imported {imported} motorela records from the 2026 workbook.")
    return imported
